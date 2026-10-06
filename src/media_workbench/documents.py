"""Editable Word and PowerPoint from one reviewed content schema.
由统一审校正文结构生成可编辑Word和PowerPoint。
"""

from pathlib import Path
import shutil
from .core import read_json, sha256


def validate(data):
    """Require explicit editorial content before generating the three document formats.
    生成三种文档格式前，要求提供明确的审校正文。
    """
    if not data.get("title") or not data.get("sections"):
        raise ValueError("Content requires title and sections / 正文必须包含title与sections")
    for section in data["sections"]:
        if not section.get("title") or not section.get("paragraphs"):
            raise ValueError("Each section requires title and paragraphs")


def build(content, output):
    """Share reviewed text across MD, Word and slides, leaving final layout for visual review.
    将审校文字同步到MD、Word与幻灯片，最终排版仍需视觉检查。
    """
    from docx import Document
    from docx.shared import Inches, Pt
    from pptx import Presentation
    from pptx.util import Inches as PInches, Pt as PPt
    from pptx.dml.color import RGBColor
    from PIL import Image

    data = read_json(content)
    validate(data)
    base = Path(content).resolve().parent
    output = Path(output)
    output.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.core_properties.author = ""
    doc.core_properties.title = data["title"]
    doc.styles["Normal"].font.size = Pt(11)
    doc.sections[0].page_width = Inches(8.27)
    doc.sections[0].page_height = Inches(11.69)
    doc.add_heading(data["title"], 0)
    deck = Presentation()
    deck.slide_width = PInches(16)
    deck.slide_height = PInches(9)
    deck.core_properties.author = ""
    deck.core_properties.title = data["title"]
    markdown = ["# " + data["title"], ""]
    for index, section in enumerate(data["sections"], 1):
        title = f"{index}. {section['title']}"
        doc.add_heading(title, 1)
        markdown += ["## " + title, ""]
        for text in section["paragraphs"]:
            doc.add_paragraph(text)
            markdown += [text, ""]
        for figure in section.get("figures", []):
            path = base / figure["file"]
            with Image.open(path) as image:
                width, height = image.size
            fit = min(6 / width, 7.5 / height)
            doc.add_picture(str(path), width=Inches(width * fit), height=Inches(height * fit))
            doc.add_paragraph(figure.get("caption", ""), style="Caption")
            # Keep the generated Markdown portable rather than linking to the input folder.
            # 生成的Markdown应可独立移动，不链接回输入目录。
            target = output / "assets" / (sha256(path)[:16] + path.suffix.lower())
            target.parent.mkdir(parents=True, exist_ok=True)
            if path.resolve() != target.resolve():
                shutil.copyfile(path, target)
            markdown += [f"![{figure.get('caption', '')}]({target.relative_to(output).as_posix()})", ""]
        # One paragraph per slide avoids compressing whole chapters into bullets.
        # 一段正文一页，避免把整章压缩为几条提纲。
        for paragraph_index, text in enumerate(section["paragraphs"]):
            if len(text) > 650:
                raise ValueError(
                    "Split a paragraph longer than 650 characters into teaching steps / 过长段落请拆成教学步骤"
                )
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
            title_box = slide.shapes.add_textbox(PInches(0.7), PInches(0.4), PInches(14.6), PInches(1))
            title_box.text_frame.text = title
            title_box.text_frame.paragraphs[0].font.size = PPt(30)
            title_box.text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 0, 0)
            figure = next(iter(section.get("figures", [])), None) if paragraph_index == 0 else None
            width = 7.5 if figure else 14.5
            box = slide.shapes.add_textbox(PInches(0.7), PInches(1.7), PInches(width), PInches(6.4))
            box.text_frame.word_wrap = True
            box.text_frame.text = text
            for p in box.text_frame.paragraphs:
                p.font.size = PPt(24)
                p.font.color.rgb = RGBColor(0, 0, 0)
            if figure:
                add_figure(slide, base / figure["file"], figure.get("caption", ""), 8.6, 2, 6.4, 5.4)
            slide.notes_slide.notes_text_frame.text = section.get("source", "")
        # Preserve all additional figures; Word and slides must not diverge silently.
        # 其他配图同样保留，不能让Word和PPT静默丢失对应内容。
        for figure in section.get("figures", [])[1:]:
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            heading = slide.shapes.add_textbox(PInches(0.7), PInches(0.4), PInches(14.6), PInches(1))
            heading.text_frame.text = title
            heading.text_frame.paragraphs[0].font.size = PPt(30)
            heading.text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 0, 0)
            add_figure(slide, base / figure["file"], figure.get("caption", ""), 0.8, 1.6, 14.4, 5.6)
            slide.notes_slide.notes_text_frame.text = section.get("source", "")
    doc.save(output / "tutorial.docx")
    deck.save(output / "tutorial.pptx")
    (output / "tutorial.md").write_text("\n".join(markdown), encoding="utf8")
    return dict(
        word="tutorial.docx",
        powerpoint="tutorial.pptx",
        markdown="tutorial.md",
        slides=len(deck.slides),
        visual_review="required: export with Word/PowerPoint or inspect manually",
    )


def add_figure(slide, path, caption, x, y, max_width, max_height):
    """Fit a picture and caption inside bounded slide space without hiding source pixels.
    在有界幻灯片区域内完整放图及图注，不裁掉原图内容。
    """
    from PIL import Image
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    if len(caption) > 250:
        raise ValueError("Split long figure captions into a teaching paragraph / 长图注请拆到教学正文")
    with Image.open(path) as image:
        width, height = image.size
    ratio = min(max_width / width, max_height / height)
    slide.shapes.add_picture(
        str(path), Inches(x), Inches(y), width=Inches(width * ratio), height=Inches(height * ratio)
    )
    box = slide.shapes.add_textbox(Inches(x), Inches(y + max_height + 0.15), Inches(max_width), Inches(0.9))
    box.text_frame.word_wrap = True
    box.text_frame.text = caption
    for paragraph in box.text_frame.paragraphs:
        paragraph.font.size = Pt(18)
        paragraph.font.color.rgb = RGBColor(0, 0, 0)
